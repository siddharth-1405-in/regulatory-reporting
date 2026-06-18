"""Export — Excel (into the SAMA template) and PDF (fpdf2) from the same outputs.

Excel: opens SAMA_CAR_SA01_Template.xlsx, writes the deterministic input and
computed VALUES into the correct cells (via the registry + repaired cell map),
overwriting the template's stale cross-sheet formulas so the exported file is
fully populated and internally consistent.

PDF: generates a CFO-ready multi-page PDF via fpdf2 (pure Python, no OS deps):
  Page 1  — Cover + Summary
  Page 2+ — Schedule 1 through Schedule 6 (one per page)

Export is not gated on sign-off; status badge in the output reflects Draft vs Signed Off.
"""
from __future__ import annotations

import io
import re

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font
from sqlalchemy.orm import Session

from ..core.config import CAR_TEMPLATE
from ..models import Narrative, ReportInstance
from ..registry.car_sa01 import elements as reg
from ..registry.car_sa01.cell_map import CELL_MAP
from . import calc_service

# fpdf2 core fonts are latin-1 only; map common unicode punctuation to ASCII and
# drop anything else so a stray character (e.g. from an AI narrative) never crashes
# the export.
_UNI_MAP = {"–": "-", "—": "-", "−": "-", "“": '"', "”": '"', "‘": "'", "’": "'",
            "…": "...", "•": "-", "→": "->", "≥": ">=", "≤": "<=", " ": " "}


def _lat1(s) -> str:
    s = str(s if s is not None else "")
    for k, v in _UNI_MAP.items():
        s = s.replace(k, v)
    return s.encode("latin-1", "replace").decode("latin-1")


def _approved_narrative_en(db: Session, instance_id: int) -> str | None:
    n = (db.query(Narrative)
         .filter(Narrative.instance_id == instance_id, Narrative.language == "en",
                 Narrative.status == "approved")
         .order_by(Narrative.version.desc()).first())
    return n.body if n and n.body else None

MONEY_FMT = "#,##0"
PCT_FMT = "0.00%"
PCT_CELLS = {("Summary", "E7"), ("Summary", "E9"), ("Summary", "E11"),
             ("Schedule 5", "E8"), ("Schedule 5", "E9"), ("Schedule 5", "E10"),
             ("Schedule 6", "D25"), ("Schedule 6", "D26"), ("Schedule 6", "D27")}

_S6_ROW = {base: row for base, _label, row in reg._S6_LINES}
_S4_ROW = {base: row for base, _label, _codes, _beta, row in reg.S4_BUSINESS_LINES}


def _row_of(cell: str) -> int:
    return int(re.search(r"\d+", cell).group())


def _set(ws, cell: str, value, fmt: str | None = None):
    c = ws[cell]
    c.value = value
    if fmt:
        c.number_format = fmt


def build_excel(db: Session, instance_id: int) -> bytes:
    inst = db.get(ReportInstance, instance_id)
    run = calc_service.latest_run(db, instance_id)
    if run is None:
        raise ValueError("no calculation run to export")
    results = run.results
    values = results["values"]
    inputs = calc_service.report_instance_service.load_inputs(db, instance_id)

    wb = load_workbook(CAR_TEMPLATE)

    # Approved executive narrative as the first sheet (top of the export)
    narrative = _approved_narrative_en(db, instance_id)
    if narrative:
        ns = wb.create_sheet("Executive Narrative", 0)
        ns["A1"] = "Executive Narrative"
        ns["A1"].font = Font(bold=True, size=14, color="492079")
        ns["A2"] = f"{inst.bank_name}  ·  {inst.period_label}  ·  Approved by the report owner"
        ns["A2"].font = Font(italic=True, color="827896")
        ns["A4"] = narrative
        ns["A4"].alignment = Alignment(wrap_text=True, vertical="top")
        ns.column_dimensions["A"].width = 120
        ns.row_dimensions[4].height = 220

    # Cover identification
    cov = wb["Cover"]
    _set(cov, "C5", inst.bank_name)
    _set(cov, "C7", inst.period_label)
    _set(cov, "C8", inst.period_end.strftime("%d/%m/%Y"))

    # 1) raw input values into their registry cells
    for code, val in inputs.items():
        e = reg.REGISTRY.get(code)
        if e and e.export_cell and e.sheet_name in wb.sheetnames:
            _set(wb[e.sheet_name], e.export_cell, float(val), MONEY_FMT)

    # 2) per-line computed results (overwrite template formulas with values)
    for name, sched in results["schedules"].items():
        ws = wb[name] if name in wb.sheetnames else None
        if ws is None:
            continue
        for line in sched["lines"]:
            code, res = line["element_code"], line["result"]
            if name == "Schedule 2":
                row = _row_of(reg.get(code).export_cell)
                if "equivalent" in line["inputs"]:
                    _set(ws, f"E{row}", line["inputs"]["equivalent"], MONEY_FMT)
                _set(ws, f"G{row}", res, MONEY_FMT)
            elif name == "Schedule 3":
                _set(ws, f"E{_row_of(reg.get(code).export_cell)}", res, MONEY_FMT)
            elif name == "Schedule 4":
                _set(ws, f"G{_S4_ROW[code]}", res, MONEY_FMT)
            elif name == "Schedule 6":
                _set(ws, f"E{_S6_ROW[code]}", res, MONEY_FMT)

    # 3) derived totals + repaired cross-sheet cells via the cell map
    for code, targets in CELL_MAP.items():
        if code not in values:
            continue
        for sheet, cell in targets:
            if sheet not in wb.sheetnames:
                continue
            fmt = PCT_FMT if (sheet, cell) in PCT_CELLS else MONEY_FMT
            _set(wb[sheet], cell, values[code], fmt)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---- PDF via fpdf2 -----------------------------------------------------------

def _money(x) -> str:
    try:
        return f"{float(x):,.0f}"
    except Exception:
        return "—"


def _pct(x) -> str:
    try:
        return f"{float(x) * 100:.2f}%"
    except Exception:
        return "—"


def build_pdf(db: Session, instance_id: int) -> tuple[bytes, str]:
    """Returns (bytes, media_type) as a real PDF using fpdf2. Falls back to HTML
    only if fpdf2 is unavailable in the environment."""
    try:
        from fpdf import FPDF
    except ImportError:
        return build_html(db, instance_id).encode("utf-8"), "text/html"

    inst = db.get(ReportInstance, instance_id)
    run = calc_service.latest_run(db, instance_id)
    if run is None:
        raise ValueError("no calculation run to export")
    m = run.results["metrics"]
    flags = run.results["flags"]
    schedules = run.results.get("schedules", {})
    values = run.results.get("values", {})

    from . import report_signoff_service
    report_status = report_signoff_service.report_status(db, instance_id)
    is_compliant = flags.get("buffer_status") == "Compliant"

    # ---- colours (RGB) -------------------------------------------------------
    PURPLE = (73, 32, 121)
    MAGENTA = (179, 30, 124)
    DARK = (34, 17, 68)
    MUTED = (130, 120, 150)
    WHITE = (255, 255, 255)
    LIGHT = (246, 242, 252)
    GREEN = (31, 138, 91)
    RED = (192, 58, 58)
    LINE = (220, 210, 235)

    pdf = FPDF(orientation="P", unit="mm", format="A4")
    pdf.set_auto_page_break(auto=True, margin=18)
    pdf.set_margins(18, 18, 18)

    def _header(label: str):
        pdf.set_fill_color(*PURPLE)
        pdf.rect(0, 0, 210, 14, "F")
        pdf.set_xy(18, 3)
        pdf.set_font("Helvetica", "B", 9)
        pdf.set_text_color(*WHITE)
        pdf.cell(0, 8, _lat1(label), ln=True)

    def _footer():
        pdf.set_y(-14)
        pdf.set_font("Helvetica", "", 7)
        pdf.set_text_color(*MUTED)
        pdf.cell(0, 5, _lat1(f"Generated by the Agentic CAR Reporting Platform  ·  {inst.bank_name}  ·  {inst.period_label}  ·  Page {pdf.page_no()}"), align="C")

    def _section_title(title: str, subtitle: str = ""):
        pdf.set_font("Helvetica", "B", 14)
        pdf.set_text_color(*DARK)
        pdf.cell(0, 8, _lat1(title), ln=True)
        if subtitle:
            pdf.set_font("Helvetica", "", 9)
            pdf.set_text_color(*MUTED)
            pdf.cell(0, 5, _lat1(subtitle), ln=True)
        pdf.set_draw_color(*MAGENTA)
        pdf.set_line_width(0.6)
        pdf.line(18, pdf.get_y(), 192, pdf.get_y())
        pdf.ln(3)

    def _kv(key: str, val: str, bold_val: bool = False):
        pdf.set_font("Helvetica", "", 9)
        pdf.set_text_color(*MUTED)
        pdf.cell(80, 6, _lat1(key))
        if bold_val:
            pdf.set_font("Helvetica", "B", 9)
        pdf.set_text_color(*DARK)
        pdf.cell(0, 6, _lat1(val), ln=True)

    def _table_header(*cols_widths):
        pdf.set_fill_color(*LIGHT)
        pdf.set_font("Helvetica", "B", 7)
        pdf.set_text_color(*PURPLE)
        for label, w in cols_widths:
            pdf.cell(w, 6, _lat1(label).upper(), border="B", fill=True)
        pdf.ln()

    def _table_row(*cells_widths, bold: bool = False, alt: bool = False):
        if alt:
            pdf.set_fill_color(*LIGHT)
        pdf.set_font("Helvetica", "B" if bold else "", 8)
        pdf.set_text_color(*DARK)
        for text, w, align in cells_widths:
            pdf.cell(w, 5.5, _lat1(text), border="B", align=align, fill=alt)
        pdf.ln()

    def _status_badge(text: str, good: bool):
        color = GREEN if good else RED
        pdf.set_fill_color(*color)
        pdf.set_text_color(*WHITE)
        pdf.set_font("Helvetica", "B", 7)
        pdf.cell(28, 5, _lat1(f"  {text}"), fill=True, border=0)
        pdf.set_text_color(*DARK)

    def _narrative_block():
        body = _approved_narrative_en(db, instance_id)
        if not body:
            return
        _section_title("Executive Narrative", "Approved by the report owner")
        pdf.set_font("Helvetica", "", 9.5)
        pdf.set_text_color(*DARK)
        pdf.multi_cell(0, 5, _lat1(body))
        pdf.ln(4)

    # ============================================================ PAGE 1: Cover + Summary
    pdf.add_page()
    _header("SAMA · Capital Adequacy Return · CAR-SA-01")

    # Uniqus branding bar
    pdf.set_fill_color(*MAGENTA)
    pdf.rect(0, 14, 5, 283, "F")

    pdf.set_xy(18, 20)
    pdf.set_font("Helvetica", "B", 22)
    pdf.set_text_color(*DARK)
    pdf.cell(0, 12, _lat1(inst.bank_name), ln=True)

    pdf.set_font("Helvetica", "", 10)
    pdf.set_text_color(*MUTED)
    pdf.cell(0, 6, _lat1(f"Capital Adequacy Return  ·  {inst.period_label}  ·  Period end {inst.period_end.strftime('%d %b %Y')}"), ln=True)
    pdf.ln(2)

    # Status badges
    pdf.set_x(18)
    _status_badge(report_status, report_status == "Signed Off")
    pdf.set_x(pdf.get_x() + 3)
    _status_badge(flags.get("buffer_status", "Unknown"), is_compliant)
    pdf.ln(8)

    # Executive narrative (only when an approved English narrative exists)
    _narrative_block()

    # Cover fields
    _section_title("Cover", "Report identification")
    for k, v in [
        ("Bank full legal name", inst.bank_name),
        ("Reporting quarter", inst.period_label),
        ("Period end", inst.period_end.strftime("%d %b %Y")),
        ("Reporting currency", getattr(inst, "currency", "SAR")),
        ("Units", getattr(inst, "units", "SAR '000")),
        ("Report status", report_status),
        ("Buffer status", flags.get("buffer_status", "—")),
    ]:
        _kv(k, str(v or "—"))

    pdf.ln(4)

    # Summary table
    _section_title("Summary", "Capital Adequacy Ratios")
    _table_header(("Metric", 75), ("Amount (SAR '000)", 45), ("Ratio", 25), ("SAMA Min", 22), ("Status", 7))

    summary_rows = [
        ("Common Equity Tier 1 (CET1)", m.get("cet1"), m.get("cet1_ratio"), 0.07),
        ("Tier 1 Capital",              m.get("tier1"), m.get("tier1_ratio"), 0.085),
        ("Total Regulatory Capital",    m.get("total_capital"), m.get("total_ratio"), 0.105),
    ]
    for label, amt, ratio, min_req in summary_rows:
        ok = (ratio or 0) >= min_req
        _table_row(
            (label, 75, "L"),
            (_money(amt), 45, "R"),
            (_pct(ratio), 25, "R"),
            (_pct(min_req), 22, "R"),
            ("Met" if ok else "Below", 7, "C"),
            bold=False,
            alt=False,
        )
    _table_row(("Total Risk-Weighted Assets", 75, "L"), (_money(m.get("total_rwa")), 99, "R"), bold=True, alt=True)

    pdf.ln(4)
    # RWA breakdown
    pdf.set_font("Helvetica", "", 8)
    pdf.set_text_color(*MUTED)
    pdf.cell(0, 5, f"Credit RWA: {_money(m.get('credit_rwa'))}    Market RWA: {_money(m.get('market_rwa'))}    Operational RWA: {_money(m.get('oprisk_rwa'))}", ln=True)

    _footer()

    # ============================================================ SCHEDULE PAGES
    SCHEDULE_PAGES = [
        ("S1", "Schedule 1", "Regulatory Capital Composition", "Schedule 1"),
        ("S2", "Schedule 2", "Credit Risk Risk-Weighted Assets", "Schedule 2"),
        ("S3", "Schedule 3", "Market Risk Capital Charge (Standardised Approach)", "Schedule 3"),
        ("S4", "Schedule 4", "Operational Risk Capital Charge (Standardised Approach)", "Schedule 4"),
        ("S5", "Schedule 5", "Capital Buffers & Combined Capital Requirement", "Schedule 5"),
        ("S6", "Schedule 6", "Reconciliation to Published Financial Statements", "Schedule 6"),
    ]

    for key, short_name, full_name, sched_key in SCHEDULE_PAGES:
        pdf.add_page()
        _header(f"CAR-SA-01  ·  {inst.bank_name}  ·  {inst.period_label}")

        pdf.set_xy(18, 20)
        _section_title(short_name, full_name)

        sched_data = schedules.get(sched_key)
        if not sched_data or not sched_data.get("lines"):
            pdf.set_font("Helvetica", "I", 9)
            pdf.set_text_color(*MUTED)
            pdf.cell(0, 6, "No data available for this schedule in the current calculation run.", ln=True)
            _footer()
            continue

        lines = sched_data["lines"]
        totals = {**sched_data.get("subtotals", {}), **sched_data.get("totals", {})}

        _table_header(("Data Element", 95), ("Result (SAR '000)", 45), ("Inputs", 34))
        for i, ln in enumerate(lines):
            inputs_str = "  ·  ".join(
                f"{k}={_pct(v) if isinstance(v, float) and 0 < v < 1 else _money(v)}"
                for k, v in (ln.get("inputs") or {}).items()
            )[:55]
            _table_row(
                (ln.get("label", ln.get("element_code", "")), 95, "L"),
                (_money(ln.get("result")), 45, "R"),
                (inputs_str, 34, "L"),
                alt=i % 2 == 0,
            )

        for k, v in totals.items():
            _table_row((k, 95, "L"), (_money(v), 79, "R"), bold=True, alt=True)

        _footer()

    # fpdf2's output() returns the document as a bytearray; wrap as bytes.
    return bytes(pdf.output()), "application/pdf"


def build_html(db: Session, instance_id: int) -> str:
    """Fallback HTML report (for debugging / preview)."""
    inst = db.get(ReportInstance, instance_id)
    run = calc_service.latest_run(db, instance_id)
    if run is None:
        raise ValueError("no calculation run to export")
    m = run.results["metrics"]
    flags = run.results["flags"]
    status_color = "#1F8A5B" if flags.get("buffer_status") == "Compliant" else "#C03A3A"
    from . import report_signoff_service
    report_status = report_signoff_service.report_status(db, instance_id)
    rs_color = "#1F8A5B" if report_status == "Signed Off" else "#492079"

    def pct(x): return f"{x*100:.2f}%"
    def money(x): return f"{x:,.0f}"

    rows = "".join(
        f"<tr><td>{label}</td><td class='num'>{val}</td></tr>"
        for label, val in [
            ("CET1 Capital", money(m["cet1"])), ("Tier 1 Capital", money(m["tier1"])),
            ("Total Regulatory Capital", money(m["total_capital"])),
            ("Credit Risk RWA", money(m["credit_rwa"])), ("Market Risk RWA", money(m["market_rwa"])),
            ("Operational Risk RWA", money(m["oprisk_rwa"])),
            ("<b>Total RWA</b>", f"<b>{money(m['total_rwa'])}</b>"),
            ("CET1 Ratio", pct(m["cet1_ratio"])), ("Tier 1 Ratio", pct(m["tier1_ratio"])),
            ("Total Capital Ratio", pct(m["total_ratio"])),
        ])

    return f"""<!doctype html><html><head><meta charset='utf-8'>
<style>
  body {{ font-family:'Helvetica',sans-serif; color:#333; margin:32px; }}
  h1 {{ color:#3B2162; font-size:22px; margin:4px 0 0; }}
  .sub {{ color:#666; font-size:12px; margin-bottom:18px; }}
  .status {{ display:inline-block; padding:4px 12px; border-radius:999px; color:#fff;
            background:{status_color}; font-size:11px; }}
  table {{ border-collapse:collapse; width:100%; margin-top:14px; font-size:12px; }}
  th,td {{ text-align:left; padding:7px 10px; border-bottom:1px solid #ECE4F2; }}
  th {{ color:#492079; font-size:10px; text-transform:uppercase; }}
  td.num {{ text-align:right; }}
</style></head><body>
  <h1>{inst.bank_name}</h1>
  <div class='sub'>{inst.period_label} · {inst.period_end:%d %b %Y}</div>
  <span class='status' style='background:{rs_color}'>Report: {report_status}</span>
  <span class='status'>Buffer: {flags.get('buffer_status')}</span>
  <table><tr><th>Metric</th><th style='text-align:right'>Value</th></tr>{rows}</table>
</body></html>"""
