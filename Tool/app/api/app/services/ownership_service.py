"""Ownership — read assignments from the Regulatory Data Ownership Matrix.

The workbook's human ownership/steward metadata is read for display and the
certification dashboard. The element->domain mapping used for certification
GATING comes from the canonical registry (`source_domain`), which is exact and
line-level; the matrix is the governance contract at the data-element level.
"""
from __future__ import annotations

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from ..core.config import OWNERSHIP_MATRIX
from ..models import OwnershipAssignment
from ..registry.car_sa01 import elements as reg

CAR_REPORT_KEY = "Capital Adequacy Return"


def parse_matrix() -> list[dict]:
    """Return CAR-relevant rows from the Data Ownership Matrix sheet."""
    wb = load_workbook(OWNERSHIP_MATRIX, data_only=True)
    ws = wb["Data Ownership Matrix"]
    rows = list(ws.iter_rows(values_only=True))
    # locate header row (contains 'Data Element')
    header_idx = next(i for i, r in enumerate(rows)
                      if r and "Data Element" in [str(c) for c in r if c])
    header = [str(c).strip() if c else "" for c in rows[header_idx]]
    col = {name: i for i, name in enumerate(header)}
    out = []
    for r in rows[header_idx + 1:]:
        if not r or not r[col.get("Regulatory Report", 0)]:
            continue
        report = str(r[col["Regulatory Report"]])
        if CAR_REPORT_KEY not in report:
            continue
        out.append({
            "data_element": str(r[col["Data Element"]] or "").strip(),
            "domain": str(r[col["Data Domain"]] or "").strip(),
            "steward": str(r[col["Domain Steward"]] or "").strip(),
            "source_system": str(r[col.get("Source System", 0)] or "").strip(),
            "frequency": str(r[col.get("Frequency", 0)] or "").strip(),
            "ai_action": str(r[col.get("AI Automation Action", 0)] or "").strip(),
        })
    return out


def seed_assignments(db: Session) -> list[OwnershipAssignment]:
    db.query(OwnershipAssignment).delete()
    created = []
    for row in parse_matrix():
        a = OwnershipAssignment(**row)
        db.add(a)
        created.append(a)
    db.flush()
    return created


def required_domains() -> set[str]:
    """Domains that own at least one input element (drives certification gating)."""
    return {e.source_domain for e in reg.INPUT_ELEMENTS if e.source_domain}


def domain_for_element(code: str) -> str | None:
    e = reg.REGISTRY.get(code)
    return e.source_domain if e else None


def elements_for_domain(domain: str) -> list[str]:
    return [e.element_code for e in reg.INPUT_ELEMENTS if e.source_domain == domain]
